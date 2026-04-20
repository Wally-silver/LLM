from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse


@dataclass
class SourceDocument:
    doc_id: str
    text: str
    source_type: str
    source_value: str
    title: str
    metadata: dict


@dataclass
class LoadResult:
    documents: list[SourceDocument]
    errors: list[dict]


class DataSourceService:
    """Load documents from multiple datasource types with unified document schema."""

    SUPPORTED_FILE_TYPES = {".txt", ".md", ".pdf", ".docx", ".html", ".htm", ".json", ".jsonl", ".csv"}

    def __init__(self, http_client, max_chars: int = 2_000_000):
        self.http_client = http_client
        self.max_chars = max_chars

    @staticmethod
    def _stable_doc_id(source_type: str, source_value: str, text: str) -> str:
        digest = hashlib.sha1(f"{source_type}|{source_value}|{text[:1000]}".encode("utf-8")).hexdigest()[:16]
        return f"{source_type}:{digest}"

    @staticmethod
    def _now_iso() -> str:
        return datetime.now(timezone.utc).isoformat()

    def _normalize_doc(self, source_type: str, source_value: str, text: str, doc_id: str | None = None, title: str | None = None, metadata: dict | None = None) -> SourceDocument:
        cleaned = (text or "").strip()
        if len(cleaned) > self.max_chars:
            cleaned = cleaned[: self.max_chars]
        meta = metadata.copy() if metadata else {}
        meta.setdefault("created_at", self._now_iso())
        meta.setdefault("chunk_source", source_value)
        return SourceDocument(
            doc_id=doc_id or self._stable_doc_id(source_type, source_value, cleaned),
            text=cleaned,
            source_type=source_type,
            source_value=source_value,
            title=title or Path(source_value).name or doc_id or "untitled",
            metadata=meta,
        )

    def _extract_file_text(self, path: Path) -> tuple[str, dict]:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore"), {"file_name": path.name, "section": "full_text"}

        if suffix in {".html", ".htm"}:
            from importlib.util import find_spec

            content = path.read_text(encoding="utf-8", errors="ignore")
            if find_spec("bs4") is not None:
                from bs4 import BeautifulSoup

                soup = BeautifulSoup(content, "html.parser")
                return soup.get_text("\n", strip=True), {"file_name": path.name, "section": "html_body"}
            return content, {"file_name": path.name, "section": "html_raw"}

        if suffix == ".pdf":
            from importlib.util import find_spec

            if find_spec("pypdf") is None:
                raise RuntimeError("pdf_extract_error: pypdf is not installed")
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages):
                pages.append(page.extract_text() or "")
            return "\n".join(pages), {"file_name": path.name, "page": "all"}

        if suffix == ".docx":
            from importlib.util import find_spec

            if find_spec("docx") is None:
                raise RuntimeError("docx_extract_error: python-docx is not installed")
            from docx import Document

            doc = Document(str(path))
            text = "\n".join([p.text for p in doc.paragraphs])
            return text, {"file_name": path.name, "section": "docx_paragraphs"}

        if suffix == ".json":
            obj = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
            return json.dumps(obj, ensure_ascii=False, indent=2), {"file_name": path.name, "section": "json"}

        if suffix == ".jsonl":
            lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
            parsed = []
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                parsed.append(json.loads(line))
            return json.dumps(parsed, ensure_ascii=False, indent=2), {"file_name": path.name, "section": "jsonl"}

        if suffix == ".csv":
            rows = []
            with path.open("r", encoding="utf-8", errors="ignore") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    rows.append(row)
            return json.dumps(rows, ensure_ascii=False, indent=2), {"file_name": path.name, "section": "csv"}

        raise ValueError(f"unsupported_file_type: {suffix}")

    async def load(self, source_type: str, source_value: str, doc_id: str | None = None) -> SourceDocument:
        result = await self.load_many(source_type=source_type, source_value=source_value, doc_id=doc_id, recursive=True)
        if result.errors and not result.documents:
            raise RuntimeError(result.errors[0]["error"])
        return result.documents[0]

    async def load_many(self, source_type: str, source_value: str, doc_id: str | None = None, recursive: bool = True) -> LoadResult:
        source_type = source_type.lower().strip()
        docs: list[SourceDocument] = []
        errors: list[dict] = []

        try:
            if source_type == "inline":
                docs.append(self._normalize_doc("inline", source_value, source_value, doc_id=doc_id, title="inline_text"))
                return LoadResult(documents=docs, errors=errors)

            if source_type == "url":
                resp = await self.http_client.get(source_value)
                resp.raise_for_status()
                parsed = urlparse(source_value)
                docs.append(self._normalize_doc("url", source_value, resp.text, doc_id=doc_id, title=parsed.netloc, metadata={"section": "url_body"}))
                return LoadResult(documents=docs, errors=errors)

            path = Path(source_value)
            if source_type in {"file", "path"}:
                if not path.exists() or not path.is_file():
                    raise FileNotFoundError(f"file_not_found: {source_value}")
                text, meta = self._extract_file_text(path)
                if not text.strip():
                    raise ValueError("empty_document")
                docs.append(self._normalize_doc("file", str(path), text, doc_id=doc_id, title=path.stem, metadata=meta))
                return LoadResult(documents=docs, errors=errors)

            if source_type in {"directory", "folder"}:
                if not path.exists() or not path.is_dir():
                    raise FileNotFoundError(f"directory_not_found: {source_value}")
                iterator = path.rglob("*") if recursive else path.glob("*")
                for fp in iterator:
                    if not fp.is_file() or fp.suffix.lower() not in self.SUPPORTED_FILE_TYPES:
                        continue
                    try:
                        text, meta = self._extract_file_text(fp)
                        if not text.strip():
                            raise ValueError("empty_document")
                        docs.append(self._normalize_doc("file", str(fp), text, title=fp.stem, metadata=meta))
                    except Exception as exc:
                        errors.append({"source": str(fp), "error": str(exc)})
                return LoadResult(documents=docs, errors=errors)

            if source_type in {"json", "jsonl", "csv"}:
                # batch records -> each record one doc
                if not path.exists() or not path.is_file():
                    raise FileNotFoundError(f"file_not_found: {source_value}")
                if source_type == "json":
                    payload = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
                    records = payload if isinstance(payload, list) else [payload]
                elif source_type == "jsonl":
                    records = [json.loads(x) for x in path.read_text(encoding="utf-8", errors="ignore").splitlines() if x.strip()]
                else:
                    with path.open("r", encoding="utf-8", errors="ignore") as f:
                        records = list(csv.DictReader(f))

                for i, r in enumerate(records):
                    text = r.get("text") if isinstance(r, dict) else str(r)
                    if not text:
                        errors.append({"source": f"{source_value}#{i}", "error": "empty_document"})
                        continue
                    rid = r.get("doc_id") if isinstance(r, dict) else None
                    title = r.get("title") if isinstance(r, dict) else None
                    meta = r.get("metadata", {}) if isinstance(r, dict) else {}
                    docs.append(self._normalize_doc(source_type, f"{source_value}#{i}", text, doc_id=rid, title=title, metadata=meta))
                return LoadResult(documents=docs, errors=errors)

            raise ValueError(f"unsupported source_type: {source_type}")
        except Exception as exc:
            errors.append({"source": source_value, "error": str(exc)})
            return LoadResult(documents=docs, errors=errors)
